"""
Пакетная обработка локальных аудиофайлов.

Очередь транскрипции локальных аудиозаписей, лежащих в
<front>/static/audio_package/<пакет>/... (та же структура, что в whisper-фронте:
пакет ггммдд_id, внутри audio/<дата>/<F_ID>.spx, опционально metadata.json).

Поток на каждую запись:
  1. Проверка на звук через ffmpeg volumedetect на СЫРОМ (без loudnorm) аудио.
     Тихие (max_volume < порога) не транскрибируются — пишется штамп [БЕЗ ЗВУКА].
  2. Запись со звуком конвертируется в wav 16к моно (highpass + loudnorm) и
     отправляется на OpenAI-совместимый backend (/v1/chat/completions, потоковый
     SSE в Gemma-диалекте: audio_url data-URI; текст — только если задан свой
     промпт). Транскрипт собирается из delta-событий; discard откатывает
     хвост, оборванный сторожем петель.
  3. Результат → <F_ID>.txt рядом с аудио; сводный <пакет>.xlsx (джойн с
     metadata.json по имени файла = F_ID) в папке пакета.

ffmpeg берётся из <front>/ffmpeg/bin. Без ffmpeg .spx обработать нельзя (ошибка),
а натив-форматы (wav/mp3/…) уходят на распознавание как есть, без проверки тишины.

Всё состояние — в <front>/local_state.db (SQLite), переживает перезапуск
сервера. Прежний local_state.json переносится в базу при первом запуске и
дальше не используется (остаётся резервной копией).
"""

import os
import re
import json
import time
import uuid
import base64
import tempfile
import threading
import datetime
import subprocess

import requests

from state_store import StateStore

# ── Пороговые значения ────────────────────────────────────────────────────────
# Запись без звука: пиковая громкость (max_volume) ниже порога в дБ.
SILENCE_MAX_DB = -40.0

AUDIO_EXTENSIONS = {'spx', 'wav', 'mp3', 'flac', 'ogg', 'm4a', 'webm'}

MIME_BY_EXT = {
    'wav': 'audio/wav', 'mp3': 'audio/mpeg', 'flac': 'audio/flac',
    'ogg': 'audio/ogg', 'm4a': 'audio/mp4', 'webm': 'audio/webm',
}

RECENT_LIMIT = 20
XLSX_MIN_INTERVAL = 15.0          # мин. интервал между перезаписями xlsx, сек

# ── Параметры ASR по умолчанию ────────────────────────────────────────────────
# repetition_penalty и timecodes не задаём: серверные умолчания (1.0 и false) —
# те, на которых мерили качество; 1.15 гасит петли, но портит слова.
ASR_DEFAULTS = {
    'temperature': 0,
    'no_repeat_ngram_size': 0,
    'loop_guard': True,
    'whole_audio': False,
}

REQUEST_TIMEOUT = 900             # сек на один файл (стрим)
MAX_STREAM_RETRIES = 3            # повторов при обрыве потока до пометки ошибкой
BACKEND_WAIT = 20                 # пауза (сек), если backend недоступен

_VOL_RE = re.compile(r'(mean|max)_volume:\s*(-?\d+(?:\.\d+)?)\s*dB')


class BackendUnavailable(Exception):
    """Backend недоступен/оборвал поток — файл не теряем, а повторяем позже.

    connect_failed=True  → не смогли даже подключиться (backend лежит);
    connect_failed=False → поток оборвался уже в процессе (возможно, backend
                           упал на конкретной записи).
    """
    def __init__(self, message, connect_failed=False):
        super().__init__(message)
        self.connect_failed = connect_failed


def clean_transcript(s):
    """Убрать <think>-рассуждения и служебные слова канала Gemma."""
    s = s or ''
    if re.search(r'</think>', s, re.I):
        s = re.sub(r'[\s\S]*</think>', '', s, flags=re.I)
    else:
        s = re.sub(r'<think>', '', s, flags=re.I)
    s = re.sub(r'</?think>', '', s, flags=re.I)
    s = re.sub(r'^\s*(?:thought|final|analysis)\b[\s:]*', '', s, flags=re.I)
    return s.strip()


class LocalProcessor:
    def __init__(self, base_dir, lm_backend=None, logger=None, resolve_backend=None,
                 list_models=None):
        self.base_dir    = base_dir
        self.lm_backend  = (lm_backend or '').rstrip('/')
        # модель → (адрес сервера, настоящее имя), когда бэкендов несколько,
        # и общий список моделей всех серверов (см. backends.py)
        self.resolve_backend = resolve_backend
        self.list_models = list_models
        self.state_file  = os.path.join(base_dir, 'local_state.json')   # только для переноса
        self.db_path     = os.path.join(base_dir, 'local_state.db')
        self.ffmpeg      = os.path.join(base_dir, 'ffmpeg', 'bin', 'ffmpeg.exe')
        self.ffprobe     = os.path.join(base_dir, 'ffmpeg', 'bin', 'ffprobe.exe')
        self.default_root = os.path.join(base_dir, 'static', 'audio_package')
        self.log = logger

        self._lock = threading.RLock()
        self._stop_event = threading.Event()
        self._thread = None
        self._meta_cache = {}         # package -> {F_ID: ticket}
        self._xlsx_last = {}          # package -> timestamp
        self._model_cache = None      # определённая модель (один раз за сессию)

        self.state = {
            'running': False,
            'started_at': None,
            'current': None,
            'recent': [],
            'files': {},
            'config': {
                'input_dir': self.default_root,
                'model': '',          # пусто → автоопределение
                'prompt': '',         # пусто → серверный DEFAULT_PROMPT
            },
        }
        self.store = StateStore(self.db_path, logger)
        self._load_state()

    # ── логирование ───────────────────────────────────────────────────────────
    def _info(self, m):
        if self.log: self.log.info(f"[local] {m}")

    def _error(self, m):
        if self.log: self.log.error(f"[local] {m}")

    # ── персистентность ───────────────────────────────────────────────────────
    # В памяти состояние живёт прежним словарём, на диск пишется построчно:
    # смена статуса одного файла трогает одну строку таблицы, а не
    # перезаписывает весь снимок, как это делал local_state.json.
    def _load_state(self):
        try:
            self.store.migrate_from_json(self.state_file, RECENT_LIMIT)
            self.state = self.store.load(self.state)
            self._info(f"Состояние загружено: {len(self.state['files'])} файлов")
        except Exception as e:
            self._error(f"Не удалось загрузить состояние: {e}")

    def _persist(self, rec=None, config=False):
        """Сохранить изменения: рантайм всегда, запись файла и настройки — по надобности."""
        self.store.save_meta({
            'started_at': self.state['started_at'],
            'running': False,          # запущенность после перезапуска не наследуется
        })
        if rec is not None:
            self.store.save_file(rec)
        if config:
            self.store.save_config(self.state['config'])

    def _push_recent(self, name, status, message=''):
        entry = {
            'name': name, 'status': status, 'message': message,
            'ts': datetime.datetime.now().isoformat(timespec='seconds'),
        }
        self.state['recent'].insert(0, entry)
        del self.state['recent'][RECENT_LIMIT:]
        self.store.push_recent(entry, RECENT_LIMIT)

    # ── конфигурация ──────────────────────────────────────────────────────────
    def set_config(self, input_dir=None, model=None, prompt=None):
        with self._lock:
            if input_dir is not None:
                self.state['config']['input_dir'] = input_dir.strip() or self.default_root
            if model is not None:
                self.state['config']['model'] = model.strip()
            if prompt is not None:
                self.state['config']['prompt'] = prompt
            if model is not None:
                self._model_cache = None   # модель сменили — забыть авто-выбор
            self._meta_cache.clear()
            self._persist(config=True)

    @property
    def audio_root(self):
        return self.state['config'].get('input_dir') or self.default_root

    def ffmpeg_available(self):
        return os.path.exists(self.ffmpeg)

    # ── сканирование ──────────────────────────────────────────────────────────
    def _packages(self):
        root = self.audio_root
        if not os.path.isdir(root):
            return []
        return [n for n in sorted(os.listdir(root))
                if os.path.isdir(os.path.join(root, n))]

    def scan(self):
        added = 0
        new_keys = []
        with self._lock:
            for package in self._packages():
                pkg_dir = os.path.join(self.audio_root, package)
                for dirpath, _dirs, files in os.walk(pkg_dir):
                    for fn in files:
                        ext = fn.rsplit('.', 1)[-1].lower() if '.' in fn else ''
                        if ext not in AUDIO_EXTENSIONS:
                            continue
                        full = os.path.join(dirpath, fn)
                        rel = os.path.relpath(full, pkg_dir).replace('\\', '/')
                        key = f"{package}/{rel}"
                        if key in self.state['files']:
                            continue
                        new_keys.append(key)
                        self.state['files'][key] = {
                            'package': package,
                            'rel_path': rel,
                            'name': fn,
                            'f_id': os.path.splitext(fn)[0],
                            'status': 'pending',
                            'txt_path': None,
                            'duration': None,
                            'max_volume_db': None,
                            'mean_volume_db': None,
                            'error': None,
                            'updated_at': None,
                        }
                        added += 1
            if added:
                # Одна транзакция на всё сканирование, а не запись на каждый файл.
                self.store.save_files([self.state['files'][k] for k in new_keys])
        self._info(f"Сканирование ({self.audio_root}): добавлено {added}")
        return added

    # ── управление ────────────────────────────────────────────────────────────
    def start(self):
        with self._lock:
            if self.state['running'] and self._thread and self._thread.is_alive():
                return False
        self.scan()
        with self._lock:
            self._stop_event.clear()
            self.state['running'] = True
            self.state['started_at'] = datetime.datetime.now().isoformat(timespec='seconds')
            self._persist()
        self._thread = threading.Thread(target=self._worker, daemon=True)
        self._thread.start()
        self._info("Обработка запущена")
        return True

    def stop(self):
        self._stop_event.set()
        with self._lock:
            self.state['running'] = False
            self.state['current'] = None
            self._persist()
        self._info("Запрошена остановка")
        return True

    # ── рабочий цикл ──────────────────────────────────────────────────────────
    def _next_pending(self):
        with self._lock:
            for key, rec in self.state['files'].items():
                if rec['status'] == 'pending':
                    rec['status'] = 'processing'
                    rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
                    self.state['current'] = {
                        'name': rec['name'], 'package': rec['package'],
                        'rel_path': rec['rel_path'], 'stage': 'analyzing',
                    }
                    self._persist(rec)
                    return key, rec
        return None, None

    def _worker(self):
        touched = set()
        try:
            while not self._stop_event.is_set():
                key, rec = self._next_pending()
                if rec is None:
                    break
                try:
                    self._process_one(key, rec)
                except BackendUnavailable as e:
                    if self._handle_unavailable(rec, e):
                        continue          # вернули в очередь — не считаем ошибкой
                    # исчерпали повторы по обрыву потока → это уже ошибка файла
                    with self._lock:
                        rec['status'] = 'error'
                        rec['error'] = f'обрыв потока: {e}'
                        rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
                        self._push_recent(rec['name'], 'error', 'обрыв потока')
                        self.state['current'] = None
                        self._persist(rec)
                    self._error(f"{rec['name']}: обрыв потока (исчерпаны повторы): {e}")
                except Exception as e:
                    with self._lock:
                        rec['status'] = 'error'
                        rec['error'] = str(e)
                        rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
                        self._push_recent(rec['name'], 'error', str(e))
                        self.state['current'] = None
                        self._persist(rec)
                    self._error(f"{rec['name']}: {e}")
                touched.add(rec['package'])
                self._maybe_write_xlsx(rec['package'])
        finally:
            for pkg in touched:
                self._write_xlsx(pkg, force=True)
            with self._lock:
                self.state['running'] = False
                self.state['current'] = None
                self._persist()
            self._info("Рабочий цикл завершён")

    def _wait(self, seconds):
        """Пауза с проверкой флага остановки (шаг 1 сек)."""
        end = time.time() + seconds
        while time.time() < end and not self._stop_event.is_set():
            time.sleep(1)

    def _handle_unavailable(self, rec, err):
        """Backend недоступен: вернуть файл в очередь. True — повторим,
        False — повторы по обрыву потока исчерпаны (пометить ошибкой)."""
        if not err.connect_failed:
            rec['stream_retries'] = rec.get('stream_retries', 0) + 1
            if rec['stream_retries'] > MAX_STREAM_RETRIES:
                return False
        with self._lock:
            rec['status'] = 'pending'          # вернём в очередь, не теряем
            rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
            self.state['current'] = None
            self._persist(rec)
        if err.connect_failed:
            self._info(f"Backend недоступен, пауза {BACKEND_WAIT}с и повтор: {err}")
        else:
            self._info(f"Обрыв потока ({rec.get('stream_retries')}/{MAX_STREAM_RETRIES}), "
                       f"повтор: {rec['name']}")
        self._wait(BACKEND_WAIT)
        return True

    def _set_stage(self, rec, stage):
        with self._lock:
            cur = self.state['current']
            if cur and cur.get('rel_path') == rec['rel_path'] \
                    and cur.get('package') == rec['package']:
                cur['stage'] = stage

    def _process_one(self, key, rec):
        pkg_dir = os.path.join(self.audio_root, rec['package'])
        src = os.path.join(pkg_dir, rec['rel_path'].replace('/', os.sep))
        txt_path = os.path.splitext(src)[0] + '.txt'
        ext = rec['name'].rsplit('.', 1)[-1].lower() if '.' in rec['name'] else ''

        tmp_base = os.path.join(tempfile.gettempdir(), 'g4_' + uuid.uuid4().hex)
        raw_wav = tmp_base + '_raw.wav'
        proc_wav = tmp_base + '_proc.wav'
        send_path, send_mime = src, MIME_BY_EXT.get(ext, 'audio/wav')

        try:
            if self.ffmpeg_available():
                # 1) анализ громкости на сыром аудио
                self._set_stage(rec, 'analyzing')
                if not self._ffmpeg_raw(src, raw_wav):
                    raise RuntimeError('ошибка конвертации для анализа')
                mean_db, max_db = self._volumedetect(raw_wav)
                duration = self._duration(raw_wav)
                with self._lock:
                    rec['duration'] = duration
                    rec['mean_volume_db'] = mean_db
                    rec['max_volume_db'] = max_db

                # 2) тишина?
                if max_db is None or max_db < SILENCE_MAX_DB:
                    with open(txt_path, 'w', encoding='utf-8') as f:
                        f.write(f"[БЕЗ ЗВУКА] max_volume={max_db} dB, "
                                f"mean_volume={mean_db} dB\n")
                    with self._lock:
                        rec['status'] = 'silent'
                        rec['txt_path'] = txt_path
                        rec['error'] = None
                        rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
                        self._push_recent(rec['name'], 'silent', f"max={max_db} dB")
                        self.state['current'] = None
                        self._persist(rec)
                    self._info(f"{rec['name']}: без звука (max={max_db} dB)")
                    return

                # 3) обработанное аудио для распознавания
                if self._ffmpeg_processed(src, proc_wav):
                    send_path, send_mime = proc_wav, 'audio/wav'
                else:
                    send_path, send_mime = raw_wav, 'audio/wav'
            else:
                # без ffmpeg: .spx декодировать нечем
                if ext == 'spx':
                    raise RuntimeError('нужен ffmpeg (front/ffmpeg/bin) для .spx')
                # натив-формат уходит как есть, без проверки тишины

            # 4) транскрипция
            self._set_stage(rec, 'transcribing')
            transcript = self._transcribe(send_path, send_mime)

            with open(txt_path, 'w', encoding='utf-8') as f:
                f.write(transcript)
            with self._lock:
                rec['status'] = 'completed'
                rec['txt_path'] = txt_path
                rec['error'] = None
                rec['updated_at'] = datetime.datetime.now().isoformat(timespec='seconds')
                self._push_recent(rec['name'], 'completed', f"{len(transcript)} симв.")
                self.state['current'] = None
                self._persist(rec)
            self._info(f"{rec['name']}: готово ({len(transcript)} симв.)")

        finally:
            for p in (raw_wav, proc_wav):
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except OSError:
                        pass

    # ── ffmpeg ────────────────────────────────────────────────────────────────
    def _ffmpeg_raw(self, src, dst):
        """Конвертация без нормализации — для честного анализа громкости."""
        cmd = [self.ffmpeg, '-i', src, '-ar', '16000', '-ac', '1',
               '-y', '-loglevel', 'error', dst]
        try:
            r = subprocess.run(cmd, capture_output=True)
            return r.returncode == 0 and os.path.exists(dst)
        except Exception:
            return False

    def _ffmpeg_processed(self, src, dst):
        """Конвертация с фильтром шума и loudnorm — для распознавания."""
        cmd = [self.ffmpeg, '-i', src,
               '-af', 'highpass=f=80,loudnorm=I=-16:TP=-1.5:LRA=11',
               '-ar', '16000', '-ac', '1', '-y', '-loglevel', 'error', dst]
        try:
            r = subprocess.run(cmd, capture_output=True)
            return r.returncode == 0 and os.path.exists(dst)
        except Exception:
            return False

    def convert_bytes_to_asr_wav(self, raw, src_ext=''):
        """Декодировать произвольные аудио-байты → 16к-моно PCM WAV байты тем же
        рецептом (highpass + loudnorm), что и локальная обработка. Нужно, чтобы
        чат-путь слал бэкенду то же, что он ожидает (бэкенд без ffmpeg).

        Возвращает WAV-байты, либо None если ffmpeg недоступен / конвертация не
        удалась (тогда вызывающий шлёт исходные байты как раньше).
        """
        if not self.ffmpeg_available():
            return None
        base = os.path.join(tempfile.gettempdir(), 'g4chat_' + uuid.uuid4().hex)
        src = base + ('.' + src_ext if src_ext else '.bin')
        dst = base + '_16k.wav'
        try:
            with open(src, 'wb') as f:
                f.write(raw)
            if self._ffmpeg_processed(src, dst):
                with open(dst, 'rb') as f:
                    return f.read()
            return None
        except Exception:
            return None
        finally:
            for p in (src, dst):
                try:
                    os.remove(p)
                except OSError:
                    pass

    def _volumedetect(self, wav):
        cmd = [self.ffmpeg, '-i', wav, '-af', 'volumedetect', '-f', 'null', '-']
        mean_db = max_db = None
        try:
            r = subprocess.run(cmd, capture_output=True)
            stderr = r.stderr.decode(errors='replace')
            for m in _VOL_RE.finditer(stderr):
                val = float(m.group(2))
                if m.group(1) == 'mean':
                    mean_db = val
                else:
                    max_db = val
        except Exception as e:
            self._error(f"volumedetect: {e}")
        return mean_db, max_db

    def _duration(self, wav):
        if not os.path.exists(self.ffprobe):
            return None
        cmd = [self.ffprobe, '-v', 'error', '-show_entries', 'format=duration',
               '-of', 'csv=p=0', wav]
        try:
            r = subprocess.run(cmd, capture_output=True, text=True)
            return round(float(r.stdout.strip()), 2)
        except Exception:
            return None

    # ── модель ────────────────────────────────────────────────────────────────
    def _resolve_model(self):
        # Приоритет: заданная в настройках → уже определённая (кэш) → авто-выбор
        # первой доступной. На сервере одна модель, поэтому определяем максимум
        # один раз за сессию; сбой определения не валит файл — реальную ошибку
        # (backend недоступен) покажет сам запрос транскрипции.
        model = self.state['config'].get('model', '').strip()
        if model:
            return model
        if self._model_cache:
            return self._model_cache
        try:
            if self.list_models:
                # Серверов несколько — берём речевую модель, а не первую попавшуюся.
                entries = self.list_models()
                ids = ([m['id'] for m in entries if m.get('kind') == 'asr']
                       or [m['id'] for m in entries])
            else:
                r = requests.get(f"{self.lm_backend}/v1/models", timeout=10)
                ids = [m['id'] for m in r.json().get('data', []) if m.get('id')]
            if ids:
                self._model_cache = ids[0]
                self._info(f"Модель определена автоматически: {ids[0]}")
                return ids[0]
        except Exception as e:
            self._error(f"Модель не определена (backend недоступен?): {e}")
        return ''   # отправим без явной модели — единственная загруженная подхватится

    # ── транскрипция через backend (SSE) ──────────────────────────────────────
    def _transcribe(self, path, mime):
        with open(path, 'rb') as f:
            b64 = base64.b64encode(f.read()).decode('ascii')

        model = self._resolve_model()
        prompt = self.state['config'].get('prompt', '') or ''
        # «qwen #2» — экземпляр на втором сервере; бэкенду нужно настоящее имя.
        base = self.lm_backend
        if self.resolve_backend:
            base, model = self.resolve_backend(model)
        if not base:
            raise BackendUnavailable('нет ни одного сервера моделей', connect_failed=True)

        # Текст рядом с аудио приклеивается к служебной просьбе модели, а её
        # учили без него, — шлём только свой промпт, если он задан.
        content = [{'type': 'audio_url', 'audio_url': {'url': f'data:{mime};base64,{b64}'}}]
        if prompt.strip():
            content.append({'type': 'text', 'text': prompt})
        payload = {
            'model': model,
            'messages': [{'role': 'user', 'content': content}],
            'stream': True,
            'chat_template_kwargs': {'enable_thinking': False},
        }
        payload.update(ASR_DEFAULTS)

        try:
            r = requests.post(f"{base}/v1/chat/completions",
                              json=payload, stream=True, timeout=REQUEST_TIMEOUT)
        except (requests.exceptions.ConnectionError,
                requests.exceptions.Timeout) as e:
            raise BackendUnavailable(str(e), connect_failed=True)

        if r.status_code != 200:
            raise RuntimeError(f"backend HTTP {r.status_code}: {r.text[:300]}")

        text = ''
        saw_error = None
        try:
            for raw in r.iter_lines(decode_unicode=True):
                if self._stop_event.is_set():
                    r.close()
                    raise RuntimeError('остановлено пользователем')
                if not raw:
                    continue
                line = raw[5:].strip() if raw.startswith('data:') else raw.strip()
                if not line:
                    continue
                if line == '[DONE]':
                    break
                try:
                    ev = json.loads(line)
                except Exception:
                    continue
                etype = ev.get('type')
                if etype == 'delta':
                    text += ev.get('text', '')
                elif etype == 'discard':
                    # сторож оборвал петлю и откатил её — убираем и у себя
                    n = int(ev.get('chars') or 0)
                    if n > 0:
                        text = text[:-n]
                elif etype == 'error':
                    saw_error = ev.get('message', 'ошибка сервера')
                    break
                elif etype is None:
                    delta = (ev.get('choices') or [{}])[0].get('delta', {})
                    text += delta.get('content', '') or ''
        except (requests.exceptions.ChunkedEncodingError,
                requests.exceptions.ConnectionError) as e:
            # поток оборвался на середине (backend упал/перезапустился)
            raise BackendUnavailable(str(e), connect_failed=False)

        if saw_error:
            raise RuntimeError(saw_error)
        return clean_transcript(text)

    # ── xlsx (на пакет, с джойном metadata.json) ──────────────────────────────
    def _load_meta(self, package):
        if package in self._meta_cache:
            return self._meta_cache[package]
        meta = {}
        path = os.path.join(self.audio_root, package, 'metadata.json')
        if os.path.exists(path):
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                for t in data.get('tickets', []):
                    meta[t.get('F_ID')] = t
            except Exception as e:
                self._error(f"metadata.json ({package}): {e}")
        self._meta_cache[package] = meta
        return meta

    def _maybe_write_xlsx(self, package):
        if time.time() - self._xlsx_last.get(package, 0) >= XLSX_MIN_INTERVAL:
            self._write_xlsx(package)

    def xlsx_path(self, package):
        return os.path.join(self.audio_root, package, f"{package}.xlsx")

    def _write_xlsx(self, package, force=False):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except Exception as e:
            self._error(f"openpyxl недоступен: {e}")
            return

        with self._lock:
            files = [r for r in self.state['files'].values()
                     if r['package'] == package]
        files.sort(key=lambda r: r['rel_path'])
        meta = self._load_meta(package)

        status_ru = {
            'completed': 'Готово', 'silent': 'Без звука', 'error': 'Ошибка',
            'pending': 'В очереди', 'processing': 'В обработке',
        }

        wb = Workbook()
        ws = wb.active
        ws.title = 'Транскрипция'
        headers = ['F_ID', 'Файл', 'Дата/время', 'Оператор', 'Очередь',
                   'Талон', 'Филиал', 'Статус', 'Длит., с', 'max dB', 'Транскрипт']
        ws.append(headers)
        hf = Font(bold=True, color='FFFFFF')
        hfill = PatternFill('solid', fgColor='4B3F8F')
        for c in ws[1]:
            c.font = hf
            c.fill = hfill
            c.alignment = Alignment(vertical='center')

        for r in files:
            t = meta.get(r['f_id'], {})
            transcript_text = ''
            if r['status'] in ('completed', 'silent') and r.get('txt_path') \
                    and os.path.exists(r['txt_path']):
                try:
                    with open(r['txt_path'], 'r', encoding='utf-8') as f:
                        transcript_text = f.read()[:32000]
                except Exception:
                    pass
            ws.append([
                r['f_id'],
                r['rel_path'],
                t.get('F_START_TIME') or '',
                t.get('F_OP_NAME') or '',
                t.get('F_Q_NAME') or '',
                t.get('F_TICKET') if t.get('F_TICKET') is not None else '',
                (t.get('branch_name') or '').strip(),
                status_ru.get(r['status'], r['status']),
                r.get('duration') if r.get('duration') is not None else '',
                r.get('max_volume_db') if r.get('max_volume_db') is not None else '',
                transcript_text,
            ])

        for i, w in enumerate([38, 40, 19, 26, 22, 8, 30, 12, 9, 8, 80], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'

        try:
            wb.save(self.xlsx_path(package))
            self._xlsx_last[package] = time.time()
        except Exception as e:
            self._error(f"Не удалось сохранить xlsx ({package}): {e}")

    # ── статус для UI ─────────────────────────────────────────────────────────
    def status_summary(self):
        with self._lock:
            files = list(self.state['files'].values())
            totals = {'total': len(files), 'completed': 0, 'silent': 0,
                      'error': 0, 'pending': 0, 'processing': 0}
            per_pkg = {}
            for r in files:
                st = r['status'] if r['status'] in totals else 'processing'
                totals[st] += 1
                pp = per_pkg.setdefault(r['package'], {
                    'name': r['package'], 'total': 0, 'completed': 0,
                    'silent': 0, 'error': 0, 'pending': 0, 'processing': 0,
                })
                pp['total'] += 1
                pp[st] += 1
            packages = []
            for _, pp in sorted(per_pkg.items()):
                pp['done'] = pp['completed'] + pp['silent']
                pp['xlsx'] = os.path.exists(self.xlsx_path(pp['name']))
                packages.append(pp)

            return {
                'running': self.state['running'],
                'started_at': self.state['started_at'],
                'current': self.state['current'],
                'totals': totals,
                'packages': packages,
                'recent': list(self.state['recent']),
                'silence_threshold_db': SILENCE_MAX_DB,
                'ffmpeg': self.ffmpeg_available(),
                'config': dict(self.state['config']),
            }
